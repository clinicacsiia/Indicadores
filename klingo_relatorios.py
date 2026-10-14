"""Motor dos paineis Faturamento e Financeiro: baixa relatorios do Klingo e acumula em Excel.

Como funciona
- Cada painel tem uma pasta com um "relatorios.json" que diz quais relatorios do Klingo usar, em que
  aba do Excel guardar e como converter as colunas (numeros e datas).
- O periodo e pedido por mes ("Livre", de/ate). Cada mes baixado e guardado com a coluna
  "Mês (filtro)": ao atualizar, o mes inteiro e SUBSTITUIDO pelo que o Klingo devolve agora. Assim
  o que mudou no Klingo (baixa de fatura, glosa, lancamento corrigido ou apagado) muda aqui tambem,
  e o que ja foi baixado de meses antigos continua guardado.
- Os dados vem da API EXTERNA OFICIAL do Klingo (api-externa.klingo.app), com o token fixo X-APP-TOKEN guardado em
  klingo_token.txt (ver extrair_laudo.ler_token). Nao ha login nem navegador. Cada resposta e conferida por
  extrair_laudo.validar_resposta antes de virar linha: o que nao passar nao entra na planilha (o mes fica como estava).
- O Excel e gravado num arquivo temporario e trocado no final: se ele estiver aberto no Excel, nada
  e perdido e o aviso pede para fechar.
"""
import json
import os
import re
import sys
import time
from calendar import monthrange
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).parent
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "Acompanhamento de META"))
import armazenamento as arm  # noqa: E402  (Supabase no lugar do Excel, quando o portal roda na Vercel)
import extrair_laudo as ex  # noqa: E402  (cliente da API do Klingo, token e config compartilhados com as Metas)
import fonte_banco as fb  # noqa: E402  (na nuvem os paineis leem o banco, que a coleta do Supabase mantem em dia)

COL_MES = "Mês (filtro)"
COL_EM = "Atualizado em"
COL_ALERTA = "Alerta"
CONTROLE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # caracteres que o Excel nao aceita
RelatorioComErro = ex.RelatorioComErro   # as mesmas excecoes das Metas (inclui RespostaInvalida)


# ---------------------------------------------------------------- definicao e arquivos
def ler_definicao(pasta: Path) -> dict:
    return json.loads((pasta / "relatorios.json").read_text(encoding="utf-8"))


def planilha_de(pasta: Path):
    if arm.ativo():   # na nuvem nao ha arquivo: o "arquivo" e o conjunto de dados do painel no Supabase
        return arm.PlanilhaVirtual(arm.painel_de_pasta(pasta))
    return pasta / ler_definicao(pasta)["arquivo"]


def mes_mais(ym: str, n: int) -> str:
    total = int(ym[:4]) * 12 + int(ym[5:]) - 1 + n
    return f"{total // 12:04d}-{total % 12 + 1:02d}"


def meses_entre(inicio: str, fim: str) -> list[str]:
    out, atual = [], inicio
    while atual <= fim:
        out.append(atual)
        atual = mes_mais(atual, 1)
    return out


# ---------------------------------------------------------------- consulta ao Klingo
def colunas_exigidas(rel: dict) -> tuple[str, ...]:
    """Colunas que o relatorios.json declara (chave, numeros, datas). Se o Klingo tirar alguma do relatorio, a resposta e recusada."""
    return tuple(dict.fromkeys([*rel.get("chave", []), *rel.get("numeros", []), *rel.get("datas", {})]))


def consultar(rel: dict, de: str, ate: str) -> dict:
    """Pede um relatorio para o periodo livre de..ate na API externa. Devolve o 'data' JA CONFERIDO (extrair_laudo.validar_resposta)."""
    filtros, periodo = {}, None
    if not rel.get("sem_periodo"):   # relatorio sem filtro de data (so "Id > 0"): vem inteiro, sem de/ate
        pid = rel["periodo"]
        filtros, periodo = {pid: "l", f"{pid}-from": de, f"{pid}-to": ate}, (pid, "l")
    return ex.consultar_relatorio(rel["id"], filtros, periodo=periodo, exigidas=colunas_exigidas(rel))


def baixar_faixa(rel: dict, de: str, ate: str, nivel: int = 0) -> list[dict]:
    """Linhas do relatorio no periodo, com o titulo real das colunas. Divide o periodo se o Klingo avisar que passou do limite."""
    d = consultar(rel, de, ate)
    if d.get("limit"):   # o Klingo cortou o resultado: so serve se for pedido em pedacos menores
        if rel.get("sem_periodo"):  # nao ha como dividir um relatorio sem filtro de data
            raise RelatorioComErro("o Klingo cortou o resultado (limite de linhas); crie um filtro de data no relatório para baixar por mês")
        if de < ate and nivel < 5:
            meio = date.fromordinal((date.fromisoformat(de).toordinal() + date.fromisoformat(ate).toordinal()) // 2)
            depois = date.fromordinal(meio.toordinal() + 1)
            return baixar_faixa(rel, de, meio.isoformat(), nivel + 1) + baixar_faixa(rel, depois.isoformat(), ate, nivel + 1)
        raise RelatorioComErro(f"o Klingo cortou o resultado (limite de linhas) mesmo para {de} a {ate}: as linhas que faltam não seriam gravadas")
    res = d["result"]
    titulos = {c["name"]: c["title"] for c in (res.get("group_cols") or []) + (res.get("columns") or [])}
    return [{titulos[k]: v for k, v in bruto.items() if k in titulos} for bruto in (res.get("data") or [])]


# ---------------------------------------------------------------- conversao de valores
def para_numero(v):
    if v in (None, ""):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return int(f) if f.is_integer() and "." not in str(v) else f


def para_data(v):
    """('2026-10-02 09:58:25' | '02/10/2026') -> (datetime, None); data absurda (ano < 1900) -> (None, texto original)."""
    if v in (None, ""):
        return None, None
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})(?:[ T](\d{2}):(\d{2})(?::(\d{2}))?)?", s)
    if m:
        a, mes, dia, h, mi, se = m.groups()
    else:
        m = re.match(r"^(\d{2})/(\d{2})/(\d{4})(?:\s+(\d{2}):(\d{2})(?::(\d{2}))?)?", s)
        if not m:
            return None, s
        dia, mes, a, h, mi, se = m.groups()
    if int(a) < 1900:
        return None, s
    try:
        return datetime(int(a), int(mes), int(dia), int(h or 0), int(mi or 0), int(se or 0)), None
    except ValueError:
        return None, s


def converter(rel: dict, brutas: list[dict], ym: str, agora: datetime) -> list[dict]:
    numeros, datas, sem_apostrofo = set(rel.get("numeros", [])), rel.get("datas", {}), set(rel.get("limpar_apostrofo", []))
    linhas = []
    for b in brutas:
        r, alertas = {}, []
        for titulo, v in b.items():
            if titulo in numeros:
                r[titulo] = para_numero(v)
            elif titulo in datas:
                r[titulo], invalida = para_data(v)
                if invalida:
                    alertas.append(f"{titulo} inválida ({invalida})")
            elif titulo in sem_apostrofo and isinstance(v, str):  # o Klingo põe ' na frente de "0-9" para o Excel não ler como data
                r[titulo] = v.lstrip("'") or None
            else:
                r[titulo] = None if v == "" else (CONTROLE.sub("", v) if isinstance(v, str) else v)
        if rel.get("alerta_datas"):
            r[COL_ALERTA] = "; ".join(alertas) or None
        r[COL_MES] = ym
        r[COL_EM] = agora
        linhas.append(r)
    return linhas


# ---------------------------------------------------------------- leitura e gravacao do Excel
def ler_abas_nuvem(pasta: Path) -> dict[str, pd.DataFrame]:
    """As abas do painel lidas do Supabase, com a mesma cara do Excel: um mes baixado de novo substitui o anterior, e quando o
    mesmo registro (`chave` do relatorios.json) aparece em dois meses vale o mais recentemente baixado."""
    defn, painel, abas = ler_definicao(pasta), arm.painel_de_pasta(pasta), {}
    for rel in defn["relatorios"]:
        linhas = arm.ler_aba(painel, rel["aba"])
        if not linhas:
            continue
        df = pd.DataFrame(linhas)
        for c in df.columns[df.dtypes == object]:   # None -> NaN, como o Excel entrega celula vazia
            df[c] = df[c].where(df[c].notna(), np.nan)
        if rel.get("chave") and COL_EM in df:
            df = df.assign(_em=pd.to_datetime(df[COL_EM], errors="coerce")).sort_values("_em", kind="stable").drop(columns="_em")
            df = df.drop_duplicates(subset=rel["chave"], keep="last")
        ordem = [COL_MES] + [c for c in rel.get("ordenar", []) if c in df]
        abas[rel["aba"]] = df.sort_values(ordem, kind="stable").reset_index(drop=True) if len(df) else df
    return abas


def ler_abas(pasta: Path) -> dict[str, pd.DataFrame]:
    if arm.ativo():
        return ler_abas_nuvem(pasta)
    caminho = planilha_de(pasta)
    if not caminho.exists():
        return {}
    return pd.read_excel(caminho, sheet_name=None)


def carregar(pasta: Path) -> dict[str, pd.DataFrame]:
    """As abas ja com numeros e datas convertidos (para os dashboards). Aba ausente -> DataFrame vazio."""
    defn = ler_definicao(pasta)
    abas = ler_abas(pasta)
    saida = {}
    for rel in defn["relatorios"]:
        df = abas.get(rel["aba"])
        if df is None:
            df = pd.DataFrame(columns=[COL_MES, COL_EM])
        for c in rel.get("numeros", []):
            if c in df:
                df[c] = pd.to_numeric(df[c], errors="coerce")
        for c in list(rel.get("datas", {})) + [COL_EM]:
            if c in df:
                df[c] = pd.to_datetime(df[c], errors="coerce")
        saida[rel["aba"]] = df
    return saida


def gravar_abas(pasta: Path, defn: dict, abas: dict[str, pd.DataFrame]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for rel in defn["relatorios"]:
        df = abas.get(rel["aba"])
        if df is None:
            continue
        ws = wb.create_sheet(rel["aba"])
        colunas = list(df.columns)
        if not colunas:  # relatorio que ainda nao devolveu nada (ex.: sem filtro no Klingo): aba vazia
            continue
        ws.append(colunas)
        formatos = dict(rel.get("datas", {}))
        formatos[COL_EM] = "dd/mm/yyyy hh:mm"
        for linha in df.itertuples(index=False):
            ws.append([None if pd.isna(v) else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v) for v in linha])
        for i, c in enumerate(colunas, 1):
            amostra = [len(str(x)) for x in df[c].head(300) if not pd.isna(x)]
            ws.column_dimensions[get_column_letter(i)].width = min(max([len(c)] + amostra[:300]) + 2, 60) if amostra else len(c) + 2
            if c in formatos:
                for celula in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    celula[0].number_format = formatos[c]
        for celula in ws[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="1F4E78")
            celula.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{len(df) + 1}"
    destino = planilha_de(pasta)
    temporario = destino.with_suffix(".tmp.xlsx")
    wb.save(temporario)
    try:
        os.replace(temporario, destino)
    except PermissionError:
        temporario.unlink(missing_ok=True)
        raise ex.PlanilhaEmUso(f"Feche o arquivo '{destino.name}' no Excel e tente de novo.")


# ---------------------------------------------------------------- atualizacao
def meses_a_baixar(rel: dict, defn: dict, df_atual: pd.DataFrame | None, hoje: date, historico: bool) -> list[str]:
    primeiro = rel.get("primeiro_mes", defn["primeiro_mes"])  # um relatorio pode comecar antes do painel (ex.: faturas emitidas em jun/2026)
    hoje_ym = hoje.strftime("%Y-%m")
    if rel.get("sem_periodo"):  # baixado inteiro de uma vez; o "mes" guardado e so o da ultima baixa
        return [hoje_ym]
    fim = mes_mais(hoje_ym, rel.get("meses_a_frente", 0))
    revisar = rel.get("revisar_meses")  # None = revisa todos os meses desde o primeiro
    inicio = primeiro if (historico or revisar is None) else max(primeiro, mes_mais(hoje_ym, -revisar))
    meses = set(meses_entre(inicio, fim))
    guardados = set(df_atual[COL_MES].dropna().unique()) if df_atual is not None and COL_MES in df_atual else set()
    meses |= {m for m in meses_entre(primeiro, hoje_ym) if m not in guardados}  # completa o que nunca foi baixado
    return sorted(meses)


def dados_estao_recentes(pasta: Path, intervalo_min: float) -> float | None:
    caminho = planilha_de(pasta)
    if not caminho.exists():
        return None
    idade = (time.time() - caminho.stat().st_mtime) / 60
    return idade if idade < intervalo_min else None


ORCAMENTO_S = float(os.environ.get("PAINEL_ORCAMENTO_S", "190"))   # a funcao da Vercel tem tempo limitado: para de baixar antes e deixa o resto para o proximo clique
PRAZO_FINAL: float | None = None   # instante (time.time) em que a atualizacao tem de parar de baixar; quem chama (portal_nuvem) fixa um so para varias atualizacoes em sequencia


def atualizar_do_banco(pasta: Path, forcar: bool = False, historico: bool = False) -> dict:
    """Atualiza o painel a partir do BANCO (coleta unica do Supabase), sem consultar o Klingo. Mesma regra de substituicao: cada
    (relatorio, mes) coberto pelo banco substitui o mes inteiro guardado no painel. Mes que o banco ainda nao coletou, relatorio
    sem filtro de periodo comprovado ou sem titulos de coluna ficam como estavam e entram em `avisos`. So regrava o mes cujo dado
    no banco e mais novo que o do painel."""
    defn, painel = ler_definicao(pasta), arm.painel_de_pasta(pasta)
    guardados = arm.meses_gravados(painel)
    hoje = date.today()
    hoje_ym = hoje.strftime("%Y-%m")
    resumo: dict[str, dict] = {}
    avisos, aguardando = [], 0

    def gravou(aba: str, ym: str, n: int) -> None:
        r = resumo.setdefault(aba, {"aba": aba, "linhas": 0, "meses": []})
        r["linhas"] += n
        r["meses"].append(ym)
        print(f">>> {aba} {ym}: {n} linhas (do banco)", flush=True)

    for rel in defn["relatorios"]:
        aba = rel["aba"]
        try:
            info = fb.relatorio(rel["id"])
            fb.conferir_pronto(info, rel["id"])
            tit = fb.titulos(info["id"])
            faltam = [t for t in colunas_exigidas(rel) if t not in tit.values()]
            if faltam:
                raise fb.DadosIncompletos(f"o relatório {rel['id']} não tem a(s) coluna(s) {', '.join(map(repr, faltam))} no banco: o Klingo mudou o relatório?")
            execs = fb.execucoes(info["id"])
        except fb.DadosIncompletos as e:
            avisos.append(f"{aba}: {e}; mantive o que já estava guardado.")
            continue

        if rel.get("sem_periodo"):   # vem inteiro: vale a coleta mais recente; o "mes" guardado e so o da ultima baixa
            if not execs:
                aguardando += 1
                continue
            ultimo = max(execs)
            quando, antes = execs[ultimo], guardados.get((aba, hoje_ym))
            if antes and quando <= antes:
                continue
            try:
                brutas = fb.brutas(info["id"], tit, ultimo, ultimo)
            except fb.DadosIncompletos as e:
                avisos.append(f"{aba}: {e}; mantive o que já estava guardado.")
                continue
            arm.gravar_mes(painel, aba, hoje_ym, converter(rel, brutas, hoje_ym, fb.hora_local(quando)), exclusivo=True)
            gravou(aba, hoje_ym, len(brutas))
            continue

        primeiro = rel.get("primeiro_mes", defn["primeiro_mes"])
        for ym in meses_entre(primeiro, mes_mais(hoje_ym, rel.get("meses_a_frente", 0))):
            quando = fb.cobertura_mes(info, ym, execs, hoje)
            if quando is None:
                if ym <= hoje_ym:
                    aguardando += 1
                continue
            antes = guardados.get((aba, ym))
            if antes and quando <= antes:
                continue   # o painel ja tem dado igual ou mais novo
            de, ate = (f"{ym}-01", f"{ym}-01") if info.get("granularidade") == "mes" else (f"{ym}-01", fb.ultimo_dia(ym))
            try:
                brutas = fb.brutas(info["id"], tit, de, ate)
            except fb.DadosIncompletos as e:
                avisos.append(f"{aba} {ym}: {e}; mantive o que já estava guardado.")
                continue
            arm.gravar_mes(painel, aba, ym, converter(rel, brutas, ym, fb.hora_local(quando)))
            gravou(aba, ym, len(brutas))

    if aguardando:
        avisos.append(f"{aguardando} mês(es) ainda não foram coletados pelo banco (a coleta roda sozinha a cada 5 minutos); ficam como estavam.")
    partes = " · ".join(f"{r['aba']} {r['linhas']:,} linhas em {len(r['meses'])} mês(es)".replace(",", ".") for r in resumo.values())
    msg = f"Atualizado a partir do banco: {partes or 'nada mais novo que o painel'}."
    if avisos:
        msg += " Atenção: " + " ".join(avisos)
    return {"atualizou": bool(resumo), "resumo": list(resumo.values()), "avisos": avisos, "parcial": bool(aguardando), "mensagem": msg}


def atualizar_nuvem(pasta: Path, forcar: bool = False, historico: bool = False, orcamento_s: float | None = None) -> dict:
    """Mesma regra da `atualizar` abaixo, gravando no Supabase. Cada (relatorio, mes) e gravado assim que baixa, entao um corte por
    tempo nao perde nada: o que ficou para tras entra no proximo clique. Ordem: o que nunca foi baixado, os 3 meses mais recentes, e
    depois os meses antigos do mais desatualizado para o mais novo (o Klingo pode corrigir mes antigo, e assim todos acabam revistos)."""
    if fb.ativa():   # padrao na nuvem: os dados vem do banco; PAINEL_FONTE=klingo volta a consultar o Klingo daqui
        return atualizar_do_banco(pasta, forcar=forcar, historico=historico)
    inicio, orcamento_s = time.time(), ORCAMENTO_S if orcamento_s is None else orcamento_s
    limite = PRAZO_FINAL or inicio + orcamento_s
    defn, painel = ler_definicao(pasta), arm.painel_de_pasta(pasta)
    minimo = ex.ler_config()["intervalo_minimo_minutos"]
    recente = None if (forcar or historico) else dados_estao_recentes(pasta, minimo)
    if recente is not None:
        msg = f"Dados atualizados há {recente:.0f} min (mínimo entre consultas: {minimo} min). Não consultei o Klingo."
        print(">>>", msg, flush=True)
        return {"atualizou": False, "mensagem": msg}

    guardados = arm.meses_gravados(painel)
    hoje, agora = date.today(), datetime.now()
    recentes = {mes_mais(hoje.strftime("%Y-%m"), -i) for i in range(3)}
    tarefas = []   # (prioridade, ordem, relatorio, mes)
    for rel in defn["relatorios"]:
        tem = {m for (a, m) in guardados if a == rel["aba"]}
        planejados = meses_a_baixar(rel, defn, pd.DataFrame({COL_MES: sorted(tem)}), hoje, historico)
        for ym in planejados:
            if ym not in tem:
                tarefas.append((0, -int(ym.replace("-", "")), rel, ym))
            elif ym in recentes:
                tarefas.append((1, -int(ym.replace("-", "")), rel, ym))
            else:
                quando = guardados.get((rel["aba"], ym))
                tarefas.append((2, quando.timestamp() if quando else 0.0, rel, ym))
    tarefas.sort(key=lambda t: (t[0], t[1]))

    resumo: dict[str, dict] = {}
    avisos, pendentes = [], 0
    for _, _, rel, ym in tarefas:
        if time.time() > limite:
            pendentes += 1
            continue
        de, ate = f"{ym}-01", f"{ym}-{monthrange(int(ym[:4]), int(ym[5:]))[1]:02d}"
        try:
            linhas = converter(rel, baixar_faixa(rel, de, ate), ym, agora)
            arm.gravar_mes(painel, rel["aba"], ym, linhas, exclusivo=bool(rel.get("sem_periodo")))
        except RelatorioComErro as e:
            avisos.append(f"{rel['aba']} {ym}: {e}; mantive o que já estava guardado.")
            continue
        r = resumo.setdefault(rel["aba"], {"aba": rel["aba"], "linhas": 0, "meses": []})
        r["linhas"] += len(linhas)
        r["meses"].append(ym)
        print(f">>> {rel['aba']} {ym}: {len(linhas)} linhas", flush=True)
    if pendentes:
        avisos.append(f"Parei antes de acabar por causa do tempo limite: {pendentes} mês(es) ficaram para o próximo clique em Atualizar.")
    partes = " · ".join(f"{r['aba']} {r['linhas']:,} linhas em {len(r['meses'])} mês(es)".replace(",", ".") for r in resumo.values())
    msg = f"Atualizado: {partes or 'nada novo'}."
    if avisos:
        msg += " Atenção: " + " ".join(avisos)
    return {"atualizou": bool(resumo), "resumo": list(resumo.values()), "avisos": avisos, "parcial": bool(pendentes), "mensagem": msg}


def atualizar(pasta: Path, forcar: bool = False, historico: bool = False) -> dict:
    """Baixa os relatorios, substitui os meses consultados e grava o Excel. Devolve {"atualizou", "mensagem", ...}.

    Erro do Klingo ou resposta fora do formato num relatorio/mes: aquele mes fica como estava e entra em "avisos".
    Token ausente ou recusado e rede fora do ar: levanta (ex.TokenInvalido, ex.FalhaDeRede) e nada e gravado."""
    if arm.ativo():
        return atualizar_nuvem(pasta, forcar=forcar, historico=historico)
    defn = ler_definicao(pasta)
    minimo = ex.ler_config()["intervalo_minimo_minutos"]
    recente = None if (forcar or historico) else dados_estao_recentes(pasta, minimo)
    if recente is not None:
        msg = f"Dados atualizados há {recente:.0f} min (mínimo entre consultas: {minimo} min). Não consultei o Klingo."
        print(">>>", msg, flush=True)
        return {"atualizou": False, "mensagem": msg}

    abas = ler_abas(pasta)
    hoje, agora = date.today(), datetime.now()
    plano = {rel["aba"]: meses_a_baixar(rel, defn, abas.get(rel["aba"]), hoje, historico) for rel in defn["relatorios"]}

    novas: dict[str, tuple[list[str], list[dict]]] = {}
    avisos: list[str] = []
    for rel in defn["relatorios"]:
        ok_meses, linhas = [], []
        for ym in plano[rel["aba"]]:
            de, ate = f"{ym}-01", f"{ym}-{monthrange(int(ym[:4]), int(ym[5:]))[1]:02d}"
            try:
                linhas += converter(rel, baixar_faixa(rel, de, ate), ym, agora)
                ok_meses.append(ym)
            except RelatorioComErro as e:
                avisos.append(f"{rel['aba']} {ym}: {e}; mantive o que já estava guardado.")
        novas[rel["aba"]] = (ok_meses, linhas)

    resumo = []
    for rel in defn["relatorios"]:
        nome = rel["aba"]
        ok_meses, linhas = novas[nome]
        antes = abas.get(nome)
        base = antes[~antes[COL_MES].isin(ok_meses)] if antes is not None and len(antes) else None
        if rel.get("sem_periodo") and ok_meses:  # veio inteiro: o que havia antes (inclusive de outros meses) sai
            base = None
        df_novo = pd.DataFrame(linhas)
        partes = [x for x in (base, df_novo) if x is not None and len(x)]
        df = pd.concat(partes, ignore_index=True) if partes else (antes if antes is not None else pd.DataFrame())
        if rel.get("chave") and len(df):
            df = df.drop_duplicates(subset=rel["chave"], keep="last")
        if len(df):
            ordem = [COL_MES] + [c for c in rel.get("ordenar", []) if c in df]
            df = df.sort_values(ordem, kind="stable").reset_index(drop=True)
        # colunas na ordem do relatorio, com as de controle no fim
        if len(df):
            fixas = [c for c in df.columns if c not in (COL_ALERTA, COL_MES, COL_EM)]
            df = df[fixas + [c for c in (COL_ALERTA, COL_MES, COL_EM) if c in df]]
        abas[nome] = df
        delta = len(df) - (len(antes) if antes is not None else 0)
        resumo.append({"aba": nome, "total": len(df), "delta": delta})
        print(f">>> {nome}: {len(df)} linhas ({delta:+d}), meses baixados: {', '.join(ok_meses) or '-'}", flush=True)

    gravar_abas(pasta, defn, abas)
    partes = " · ".join(f"{r['aba']} {r['total']:,}".replace(",", ".") + (f" ({r['delta']:+d})" if r["delta"] else "") for r in resumo)
    msg = f"Atualizado: {partes}."
    if avisos:
        msg += " Atenção: " + " ".join(avisos)
    return {"atualizou": True, "resumo": resumo, "avisos": avisos, "mensagem": msg}
