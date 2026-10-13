"""Extrai o relatório de Acompanhamento de Metas do Klingo e atualiza a planilha e o dashboard.

Como funciona
- Os dados ficam acumulados na aba "Dados" de "Acompanhamento de META.xlsx", sem duplicar
  (chave: coluna "Exames Realizados", que é o ID do atendimento).
- Os dados vêm da API EXTERNA OFICIAL do Klingo (api-externa.klingo.app), com um token fixo no cabeçalho
  X-APP-TOKEN: não há login, navegador nem token que vence. O token fica numa linha do arquivo
  klingo_token.txt (ver ler_token) ou na variável KLINGO_APP_TOKEN, nunca dentro do código.
- Toda resposta da API é conferida (validar_resposta) ANTES de chegar na planilha: se o formato fugir do esperado
  (relatório errado, período não aplicado, coluna que sumiu, valor de tipo estranho, linha sem ID ou sem data),
  nada é gravado e o motivo aparece na tela.
- O período é pedido explicitamente (Mês Atual), sem depender do filtro salvo no Klingo.
- Se a planilha foi atualizada há menos de "intervalo_minimo_minutos" (config.json), não consulta
  de novo: só regenera o dashboard.

Uso
  python extrair_laudo.py              atualiza se os dados estiverem velhos; sempre regenera o dashboard
  python extrair_laudo.py --forcar     consulta agora, ignorando o intervalo mínimo (o "botão")
  python extrair_laudo.py --historico  baixa o ano atual inteiro (use uma vez para completar o histórico)
  python extrair_laudo.py --sem-abrir  não abre o dashboard no navegador no final

Requisitos: pip install openpyxl pandas
"""
import gzip
import http.client
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import webbrowser
import zlib
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PASTA = Path(__file__).parent
URL = (PASTA / "Link dos dados das meta relatorio.txt").read_text(encoding="utf-8").strip()
ID_RELATORIO = int(re.search(r"/relatorio/(\d+)", URL).group(1))


def _pasta_dados_locais() -> Path:
    """Onde fica o token da API (klingo_token.txt): fora do projeto, em qualquer sistema (Windows, macOS, Linux)."""
    if os.environ.get("LOCALAPPDATA"):
        base = Path(os.environ["LOCALAPPDATA"])
    elif sys.platform == "darwin":
        base = Path.home() / "Library" / "Application Support"
    else:
        base = Path(os.environ.get("XDG_DATA_HOME") or Path.home() / ".local" / "share")
    return base / "ExtratorLaudo"


DADOS_LOCAIS = _pasta_dados_locais()
CONFIG = PASTA / "config.json"
PLANILHA = PASTA / "Acompanhamento de META.xlsx"
ABA = "Dados"
CHAVE = "Exames Realizados"
COLUNAS = [CHAVE, "Unidade", "Procedimento - Classes", "Especialidade",
           "Data do Atendimento", "Procedimento", "Mês", "Convênio"]
COLUNAS_API = COLUNAS[:-2]   # as que o relatório do Klingo traz e são obrigatórias; "Mês" é calculada aqui
# "Convênio" (no Klingo a coluna se chama "Operadora") é opcional: entra na planilha quando existe no relatório #13339 do Klingo (separa SUS de não-SUS nas metas)
LARGURAS = [18, 34, 22, 20, 20, 62, 10, 28]
API_EXTERNA = "https://api-externa.klingo.app/api/info/relatorio"
ARQUIVO_TOKEN = "klingo_token.txt"
VARIAVEL_TOKEN = "KLINGO_APP_TOKEN"
TIMEOUT_API = 180        # segundos por tentativa
TENTATIVAS_API = 3       # rede fora do ar, erro 5xx e limite de requisições (429) são repetidos; token recusado é erro do Klingo, não
ESPERA_TENTATIVA = 2     # segundos x número da tentativa, entre uma tentativa e outra
COTA = {"limite": None, "restante": None}   # última cota de requisições informada pelo Klingo (diagnóstico)
ESCALARES = (str, int, float, bool, type(None))   # o que pode haver numa célula do relatório


# ---------------------------------------------------------------- erros
class ErroApiKlingo(RuntimeError):
    """Base de tudo que pode dar errado ao buscar um relatório na API do Klingo."""


class TokenInvalido(ErroApiKlingo):
    """Token ausente ou recusado: só o usuário resolve (repetir não adianta)."""


class FalhaDeRede(ErroApiKlingo):
    """Sem conexão, tempo esgotado ou servidor do Klingo fora do ar, mesmo depois de repetir."""


class RelatorioComErro(ErroApiKlingo):
    """O Klingo respondeu com erro para este relatório ou período."""


class RespostaInvalida(RelatorioComErro):
    """A resposta não passou na conferência: estrutura, colunas, tipos ou campos obrigatórios fora do esperado."""


class PlanilhaEmUso(RuntimeError):
    """A planilha está aberta no Excel e não pode ser gravada."""


# ---------------------------------------------------------------- configuração e token
def ler_config() -> dict:
    padrao = {"intervalo_minimo_minutos": 30}
    if CONFIG.exists():
        padrao.update(json.loads(CONFIG.read_text(encoding="utf-8-sig")))
    return padrao


def locais_do_token() -> list[Path]:
    """Onde procurar o klingo_token.txt: 1) pasta de dados do usuário, fora do projeto; 2) ao lado do portal
    (um .exe numa pasta compartilhada serve a todos os computadores de uma vez)."""
    return [DADOS_LOCAIS / ARQUIVO_TOKEN, PASTA.parent / ARQUIVO_TOKEN]


def _primeira_linha_util(caminho: Path) -> str | None:
    try:
        for linha in caminho.read_text(encoding="utf-8-sig").splitlines():
            linha = linha.strip()
            if linha and not linha.startswith("#"):
                return linha
    except OSError:
        pass
    return None


def ler_token() -> str:
    """Token da API externa: variável KLINGO_APP_TOKEN ou primeira linha útil do klingo_token.txt. Nunca vai para log nem para mensagem."""
    token = os.environ.get(VARIAVEL_TOKEN, "").strip() or next((t for t in map(_primeira_linha_util, locais_do_token()) if t), "")
    if not token:
        raise TokenInvalido(f"Falta o token da API do Klingo. Salve-o (uma linha só) em {locais_do_token()[0]} "
                            f"ou defina a variável de ambiente {VARIAVEL_TOKEN}.")
    if not re.fullmatch(r"[\x21-\x7e]{8,200}", token):   # ASCII visível, sem espaço nem quebra de linha: um \n no cabeçalho derrubaria o pedido
        raise TokenInvalido(f"O token da API do Klingo está em formato inválido: deve ser uma linha só, sem espaços (confira {locais_do_token()[0]}).")
    return token


# ---------------------------------------------------------------- API externa: pedido
def _corpo_do_erro(e: urllib.error.HTTPError) -> str:
    try:
        corpo = e.read()
        if e.headers.get("Content-Encoding") == "gzip":
            corpo = gzip.decompress(corpo)
        return re.sub(r"\s+", " ", corpo.decode("utf-8", "replace")).strip()[:160]
    except Exception:
        return ""


def _baixar(url: str, token: str) -> bytes:
    """GET com X-APP-TOKEN, gzip e tentativas. Devolve o corpo já descomprimido."""
    req = urllib.request.Request(url, headers={
        "X-APP-TOKEN": token, "Accept": "application/json", "Accept-Encoding": "gzip", "User-Agent": "PortalCSI/1.0"})
    motivo = ""
    for tentativa in range(1, TENTATIVAS_API + 1):
        espera = ESPERA_TENTATIVA * tentativa
        try:
            with urllib.request.urlopen(req, timeout=TIMEOUT_API) as resp:
                COTA.update(limite=resp.headers.get("X-RateLimit-Limit"), restante=resp.headers.get("X-RateLimit-Remaining"))
                corpo = resp.read()
                if resp.headers.get("Content-Encoding") == "gzip":
                    corpo = gzip.decompress(corpo)
                return corpo
        except urllib.error.HTTPError as e:
            texto = _corpo_do_erro(e)
            if e.code in (401, 403):
                raise TokenInvalido(f"O Klingo recusou o token da API (HTTP {e.code}: {texto or 'sem detalhe'}). "
                                    f"Confira a linha do arquivo {locais_do_token()[0]}.") from None
            if e.code == 429:
                try:
                    espera = min(max(int(e.headers.get("Retry-After", espera)), 1), 60)
                except ValueError:
                    pass
                motivo = "limite de requisições do Klingo atingido (HTTP 429)"
            elif e.code >= 500:
                motivo = f"o servidor do Klingo falhou (HTTP {e.code})"
            else:
                raise RelatorioComErro(f"o Klingo recusou o pedido (HTTP {e.code}: {texto or 'sem detalhe'})") from None
        except (OSError, http.client.HTTPException, EOFError, zlib.error) as e:   # sem rede, DNS, tempo esgotado, conexão cortada, gzip truncado
            motivo = f"sem resposta do Klingo ({type(e).__name__}: {getattr(e, 'reason', None) or e})"
        if tentativa < TENTATIVAS_API:
            time.sleep(espera)
    raise FalhaDeRede(f"Não consegui falar com a API do Klingo depois de {TENTATIVAS_API} tentativas: {motivo}.")


def consultar_relatorio(id_relatorio: int, filtros: dict | None = None, periodo: tuple | None = None, exigidas: tuple = ()) -> dict:
    """Pede um relatório a API externa e devolve o 'data' JÁ CONFERIDO, no formato que o fluxo principal consome:
    {"id", "title", "inputs", "result": {"columns", "rows", "data", ["group_cols"]}, "limit", "error", ...}.

    filtros:  {"periodo0": "l", "periodo0-from": "2026-10-01", "periodo0-to": "2026-10-07"}  (o ':' que o Klingo usa nos nomes é posto aqui)
    periodo:  (id do filtro, código) só para conferir se o Klingo aplicou mesmo o período pedido, ex. ("periodo0", "l")
    exigidas: títulos de colunas que o painel precisa; se faltar alguma, a resposta é recusada
    """
    token = ler_token()
    params = {"id": str(id_relatorio), **{f":{k}": str(v) for k, v in (filtros or {}).items()}}
    corpo = _baixar(f"{API_EXTERNA}?{urllib.parse.urlencode(params)}", token)
    try:
        payload = json.loads(corpo)
    except ValueError:
        inicio = re.sub(r"\s+", " ", corpo[:80].decode("utf-8", "replace"))
        raise RespostaInvalida(f"a resposta do Klingo não é um JSON válido (começa com: {inicio!r})") from None
    return normalizar_celulas(validar_resposta(payload, id_relatorio, periodo=periodo, exigidas=exigidas))


def normalizar_celulas(payload: dict) -> dict:
    """Deixa as células do jeito que o fluxo antigo as entregava: texto (ou vazio).

    O Klingo alterna, entre uma chamada e outra IGUAL, o mesmo campo entre número (1) e texto ("1"): depende de qual servidor atende.
    Sem isto, colunas como "Parcelas" ou "Nº NF" ficariam ora número, ora texto na planilha. As colunas numéricas de verdade
    (lista "numeros" do relatorios.json) continuam sendo convertidas depois, por para_numero, que aceita as duas formas."""
    for linha in payload["result"]["data"]:
        for campo, valor in linha.items():
            if type(valor) in (int, float):   # não mexe em bool (true/false) nem em vazio
                linha[campo] = str(valor)
    return payload


# ---------------------------------------------------------------- API externa: conferência da resposta
def validar_resposta(payload, id_relatorio: int, periodo: tuple | None = None, exigidas: tuple = ()) -> dict:
    """Confere a resposta ANTES de repassá-la ao fluxo principal. Devolve o mesmo 'data' (com result sempre em forma de dict,
    mesmo quando o período não tem nenhum registro) ou levanta RelatorioComErro / RespostaInvalida. Não grava nem muda nada."""
    if not isinstance(payload, dict):
        raise RespostaInvalida(f"esperava um objeto JSON e veio {type(payload).__name__}")
    if "id" not in payload and "result" not in payload:   # o Klingo avisa erro com HTTP 200: {"status_code": 400, "msg": "Relatório não encontrado: 123"}
        msg = payload.get("msg") or payload.get("message") or payload.get("error")
        raise RelatorioComErro(f"o Klingo devolveu erro: {msg}" if msg else f"a resposta não é um relatório (chaves: {', '.join(map(str, payload))[:100] or 'nenhuma'})")
    if payload.get("error"):   # ex.: filtro de período recusado vem como error "."
        raise RelatorioComErro(f"o Klingo devolveu erro para o relatório {id_relatorio} (erro {payload['error']!r}): confira o período e os filtros")
    try:
        recebido = int(payload.get("id"))
    except (TypeError, ValueError):
        recebido = None
    if recebido != int(id_relatorio):
        raise RespostaInvalida(f"pedi o relatório {id_relatorio} e o Klingo devolveu o {payload.get('id')!r}")

    resultado = payload.get("result")
    if periodo:   # chave de filtro errada = o Klingo ignora e devolve o período padrão (hoje), sem erro nenhum
        chave, codigo = periodo
        filtros = [i for i in payload.get("inputs") or [] if isinstance(i, dict)]
        achado = next((i for i in filtros if i.get("id") == chave), None)
        if achado is None:
            de_periodo = [str(i.get("id")) for i in filtros if i.get("type") == "per-date"]
            # Sem nenhum filtro de data E sem linhas, não há o que separar por período (ex.: #13375 ainda sem glosas): vale como vazio.
            # Com linhas, ou com filtro de data de outro nome, o dado cairia no período errado ou repetido em cada mês: recusa.
            if de_periodo or resultado:
                raise RespostaInvalida(f"o relatório {id_relatorio} não tem o filtro de período {chave!r} (filtro(s) de data dele: {', '.join(de_periodo) or 'nenhum'}; "
                                       f"todos: {', '.join(str(i.get('id')) for i in filtros) or 'nenhum'}); confira \"periodo\" ou \"sem_periodo\" no relatorios.json")
        elif achado.get("value") != codigo:
            raise RespostaInvalida(f"o Klingo aplicou o período {achado.get('value')!r} em vez de {codigo!r} no relatório {id_relatorio}")

    if not resultado:   # período sem nenhum registro: o Klingo manda "" (ou nada)
        return {**payload, "result": {"columns": [], "rows": 0, "data": []}}
    if not isinstance(resultado, dict):
        raise RespostaInvalida(f"result veio como {type(resultado).__name__}; esperava um objeto")
    colunas, grupos, dados = resultado.get("columns"), resultado.get("group_cols") or [], resultado.get("data")
    if not isinstance(colunas, list) or not colunas:
        raise RespostaInvalida("result.columns ausente ou vazio")
    if not isinstance(grupos, list) or not isinstance(dados, list):
        raise RespostaInvalida("result.group_cols e result.data devem ser listas")

    titulos: dict[str, str] = {}   # name -> title, como o fluxo principal usa
    for c in grupos + colunas:
        if not isinstance(c, dict) or not isinstance(c.get("name"), str) or not c["name"] \
                or not isinstance(c.get("title"), str) or not c["title"].strip():
            raise RespostaInvalida(f"coluna sem name/title: {str(c)[:100]}")
        if c["name"] in titulos:
            raise RespostaInvalida(f"coluna repetida: {c['name']!r}")
        if c["title"] in titulos.values():   # o fluxo monta cada linha por título: duas colunas com o mesmo se sobrescreveriam
            raise RespostaInvalida(f"duas colunas com o mesmo título: {c['title']!r}")
        titulos[c["name"]] = c["title"]
    for n, linha in enumerate(dados, 1):
        if not isinstance(linha, dict):
            raise RespostaInvalida(f"linha {n} não é um objeto ({type(linha).__name__})")
        for k, v in linha.items():
            if k not in titulos:
                raise RespostaInvalida(f"linha {n} tem a chave {k!r}, que não é coluna do relatório")
            if not isinstance(v, ESCALARES):
                raise RespostaInvalida(f"linha {n}, coluna {titulos[k]!r}: valor do tipo {type(v).__name__} (esperava texto, número ou vazio)")
    if not payload.get("limit"):   # com "limit" o Klingo avisa que cortou o resultado: quem pede decide (dividir o período)
        try:
            anunciadas = int(resultado.get("rows"))
        except (TypeError, ValueError):
            raise RespostaInvalida(f"result.rows inválido: {resultado.get('rows')!r}") from None
        if anunciadas != len(dados):
            raise RespostaInvalida(f"o Klingo anunciou {anunciadas} linhas e mandou {len(dados)} (resposta incompleta)")
    faltam = [t for t in exigidas if t not in titulos.values()]
    if faltam:
        raise RespostaInvalida(f"o relatório {id_relatorio} não trouxe a(s) coluna(s) {', '.join(map(repr, faltam))}: o Klingo mudou o relatório?")
    return payload


NOME_PERIODO = {"m": "mês atual", "ma": "mês anterior", "a": "ano atual"}


def buscar_dados(periodos: list[str]) -> list[dict]:
    """Consulta cada período (m = Mês Atual, ma = Mês Anterior, a = Ano Atual) e devolve o 'data' conferido de cada um."""
    return [consultar_relatorio(ID_RELATORIO, {"periodo0": p}, periodo=("periodo0", p), exigidas=tuple(COLUNAS_API)) for p in periodos]


# ---------------------------------------------------------------- planilha
def para_linhas(data: dict) -> list[dict]:
    """Converte o JSON do Klingo em linhas com os títulos reais das colunas.

    Recusa a resposta inteira se alguma linha não tiver ID inteiro ou data válida (o que quebraria a planilha)."""
    titulos = {c["name"]: c["title"] for c in data["result"]["columns"]}
    linhas, defeitos = [], []
    for n, bruto in enumerate(data["result"]["data"], 1):
        r = {titulos[k]: v for k, v in bruto.items() if k in titulos}
        classes = r.get("Procedimento - Classes") or ""
        try:
            classes = ", ".join(json.loads(classes))
        except (ValueError, TypeError):
            pass
        try:
            ident = int(r[CHAVE])
            quando = datetime.fromisoformat(r["Data do Atendimento"])
        except (KeyError, TypeError, ValueError):
            defeitos.append(f"linha {n} (ID {r.get(CHAVE)!r}, data {r.get('Data do Atendimento')!r})")
            continue
        linhas.append({
            CHAVE: ident,
            "Unidade": r.get("Unidade"),
            "Procedimento - Classes": classes,
            "Especialidade": r.get("Especialidade"),
            "Data do Atendimento": quando,
            "Procedimento": r.get("Procedimento"),
            "Mês": quando.strftime("%Y-%m"),
            "Convênio": (r.get("Operadora") or r.get("Convênio") or "").strip() or None,
        })
    if defeitos:
        raise RespostaInvalida(f"{len(defeitos)} linha(s) sem ID ou sem data válida, nada foi gravado: " + "; ".join(defeitos[:3]) + ("…" if len(defeitos) > 3 else ""))
    return linhas


def ler_existentes(ws) -> dict[int, dict]:
    cab = [c.value for c in ws[1]]
    existentes = {}
    for valores in ws.iter_rows(min_row=2, values_only=True):
        linha = dict(zip(cab, valores))
        if linha.get(CHAVE) is not None:
            existentes[int(linha[CHAVE])] = {c: linha.get(c) for c in COLUNAS}
    return existentes


def preparar_aba():
    if PLANILHA.exists():
        wb = load_workbook(PLANILHA)
        if ABA in wb.sheetnames:
            return wb, wb[ABA], False
        ws = wb.create_sheet(ABA)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = ABA
    return wb, ws, True


def gravar(linhas: list[dict]):
    wb, ws, nova = preparar_aba()
    existentes = {} if nova else ler_existentes(ws)
    antes = set(existentes)
    for l in linhas:
        existentes[l[CHAVE]] = l  # dado novo prevalece sobre o já gravado
    todas = sorted(existentes.values(), key=lambda l: (l["Data do Atendimento"], l[CHAVE]))

    if nova:
        ws.append(COLUNAS)
        for i, larg in enumerate(LARGURAS, 1):
            ws.column_dimensions[get_column_letter(i)].width = larg
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
    else:
        ws.delete_rows(2, ws.max_row)
        for i, nome in enumerate(COLUNAS, 1):   # planilha antiga sem a coluna Convênio: completa o cabeçalho
            if ws.cell(1, i).value != nome:
                ws.cell(1, i, nome)
                ws.column_dimensions[get_column_letter(i)].width = LARGURAS[i - 1]
                ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
                ws.cell(1, i).fill = PatternFill("solid", fgColor="1F4E78")
                ws.cell(1, i).alignment = Alignment(horizontal="center", vertical="center")

    for l in todas:
        ws.append([l[c] for c in COLUNAS])
    for linha in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        linha[0].number_format = "dd/mm/yyyy hh:mm"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{len(todas) + 1}"

    try:
        wb.save(PLANILHA)
    except PermissionError:
        raise PlanilhaEmUso(f"Feche o arquivo '{PLANILHA.name}' no Excel e tente de novo.")

    novos = len(set(existentes) - antes)
    return novos, len(todas)


# ---------------------------------------------------------------- principal
def dados_estao_recentes(intervalo_min: float) -> float | None:
    """Minutos desde a última atualização da planilha, se for menor que o intervalo mínimo."""
    if not PLANILHA.exists():
        return None
    idade = (time.time() - PLANILHA.stat().st_mtime) / 60
    return idade if idade < intervalo_min else None


def atualizar_nuvem(forcar: bool = False, historico: bool = False) -> dict:
    """Mesma consulta da `atualizar`, gravando no Supabase (um mes por vez, substituindo o mes inteiro) em vez de no Excel."""
    sys.path.insert(0, str(PASTA.parent))
    import armazenamento as arm
    periodos = ["a"] if historico else ["m"] + (["ma"] if date.today().day <= 5 else [])
    minimo = ler_config()["intervalo_minimo_minutos"]
    versao = arm.versao_dados(arm.PAINEL_METAS)
    if not (forcar or historico) and versao and (time.time() - versao) / 60 < minimo:
        msg = f"Dados atualizados há {(time.time() - versao) / 60:.0f} min (mínimo entre consultas: {minimo} min). Não consultei o Klingo."
        return {"atualizou": False, "mensagem": msg}
    por_mes: dict[str, list[dict]] = {}
    for periodo, data in zip(periodos, buscar_dados(periodos)):
        print(f">>> {NOME_PERIODO.get(periodo, periodo)}: {data['result']['rows']} linhas")
        for linha in para_linhas(data):
            por_mes.setdefault(linha["Mês"], []).append(linha)
    total = sum(arm.gravar_mes(arm.PAINEL_METAS, arm.ABA_METAS, mes, linhas) for mes, linhas in sorted(por_mes.items()))
    return {"atualizou": True, "novos": total, "total": total, "mensagem": f"Metas atualizadas: {total} atendimento(s) em {len(por_mes)} mês(es)."}


def atualizar(forcar: bool = False, historico: bool = False) -> dict:
    """Consulta o Klingo (se os dados estiverem velhos ou se forcar) e grava na planilha.

    Devolve {"atualizou": bool, "mensagem": str, ...}. Usado pelo extrator e pelo botão do painel.
    """
    if os.environ.get("VERCEL") or os.environ.get("PORTAL_ARMAZENAMENTO", "").lower() == "supabase":
        return atualizar_nuvem(forcar=forcar, historico=historico)
    if historico:
        periodos = ["a"]
    else:
        periodos = ["m"] + (["ma"] if date.today().day <= 5 else [])  # nos primeiros dias, fecha o mês anterior

    minimo = ler_config()["intervalo_minimo_minutos"]
    recente = None if (forcar or historico) else dados_estao_recentes(minimo)
    if recente is not None:
        msg = f"Dados atualizados há {recente:.0f} min (mínimo entre consultas: {minimo} min). Não consultei o Klingo."
        print(">>>", msg, "Use --forcar para atualizar agora.")
        return {"atualizou": False, "mensagem": msg}

    linhas = []
    for periodo, data in zip(periodos, buscar_dados(periodos)):
        print(f">>> {NOME_PERIODO.get(periodo, periodo)}: {data['result']['rows']} linhas")
        linhas += para_linhas(data)
    novos, total = gravar(linhas)
    print(f">>> {PLANILHA.name}: {novos} novos, {total} no total.")
    return {"atualizou": True, "novos": novos, "total": total,
            "mensagem": f"Atualizado: {novos} atendimento(s) novo(s), {total} no total."}


def main():
    args = set(sys.argv[1:])
    try:
        atualizar(forcar="--forcar" in args, historico="--historico" in args)
    except PlanilhaEmUso as ex:
        raise SystemExit(str(ex))
    except ErroApiKlingo as ex:
        raise SystemExit(f"Não consegui atualizar: {ex}")

    from gerar_dashboard import gerar
    dashboard = gerar()
    print(f">>> Dashboard atualizado: {dashboard.name}")
    if "--sem-abrir" not in args:
        webbrowser.open(dashboard.as_uri())


if __name__ == "__main__":
    main()
