"""Lê a planilha mensal de despesas (aba "Planilha2" do FORNECEDORES - MM.AAAA.xlsx) e grava Despesas/planilha_base.json.

É o que aparece na aba "Planilha" do painel de Despesas: grupos (IMPOSTOS, PESSOAL...), itens e os valores de cada mês de ANTES do Klingo
(o painel usa o Klingo a partir de `inicio_confiavel`). O JSON é a base do programa local e a origem das migrations de seed; no site os
mesmos dados ficam nas tabelas despesa_* do Supabase (editáveis no Table Editor).

Uso:  python Despesas/importar_planilha.py "FORNECEDORES - 09.2026.xlsx" ["mapeamento_nomenclatura_despesas.xlsx"]
Confere se cada total de grupo da planilha bate com a soma dos itens e avisa o que não bater (nada é corrigido às escondidas).
"""
import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

SAIDA = Path(__file__).parent / "planilha_base.json"
GRUPOS = ["IMPOSTOS", "PESSOAL", "SERVIÇOS DE TERCEIROS", "DESPESAS OPERACIONAIS", "BENS PATRIMONIAIS", "DESPESAS FINANCEIRAS",
          "OUTROS IMP. CONTR. E TAXA", "EMPREST.E FINANC. AMORTIZ", "MATERIAS E MEDICAMENTOS", "FUNDO FIXO"]
COL_PRIMEIRO_MES, COL_ULTIMO_MES, COL_TOTAL = 2, 12, 13   # B..L = Out/25..Ago/26; M = "Valor Total"
PRIMEIRO_MES = (2025, 10)


def meses_da_planilha() -> list[str]:
    ano, mes, saida = PRIMEIRO_MES[0], PRIMEIRO_MES[1], []
    for _ in range(COL_ULTIMO_MES - COL_PRIMEIRO_MES + 1):
        saida.append(f"{ano:04d}-{mes:02d}")
        ano, mes = (ano + 1, 1) if mes == 12 else (ano, mes + 1)
    return saida


def conferir_cabecalho(ws, meses: list[str]) -> None:
    """Os rótulos dos meses vêm misturados (texto "Out/25" e datas); confere que a ordem é a esperada em vez de adivinhar."""
    abrev = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
    for col, ym in zip(range(COL_PRIMEIRO_MES, COL_ULTIMO_MES + 1), meses):
        v = ws.cell(1, col).value
        if isinstance(v, datetime):
            ok = v.strftime("%Y-%m") == ym
        else:
            ok = str(v).strip().lower() == f"{abrev[int(ym[5:]) - 1]}/{ym[2:4]}".lower()
        if not ok:
            raise SystemExit(f"Cabeçalho inesperado na coluna {col}: {v!r} (esperava {ym}). Confira a planilha antes de importar.")


def ler(caminho: Path) -> dict:
    ws = openpyxl.load_workbook(caminho, data_only=True)["Planilha2"]
    meses = meses_da_planilha()
    conferir_cabecalho(ws, meses)
    grupos, avisos, grupo, fechado = [], [], None, True
    totais = {}
    for r in range(2, ws.max_row + 1):
        nome = ws.cell(r, 1).value
        if not nome:
            continue
        nome = str(nome).strip()
        valores = [ws.cell(r, c).value for c in range(COL_PRIMEIRO_MES, COL_ULTIMO_MES + 1)]
        if nome == "Total":
            totais[grupo["grupo"]] = valores
            fechado = True
            continue
        if nome == "Total das despesas":
            totais["_geral"] = valores
            break
        if fechado and nome in GRUPOS:
            grupo = {"grupo": nome, "itens": []}
            grupos.append(grupo)
            fechado = False
            if any(v is not None for v in valores):
                avisos.append(f"A linha do grupo {nome} traz valores soltos ({[v for v in valores if v is not None]}) que o total do grupo não inclui; ignorados.")
            continue
        grupo["itens"].append({"item": nome, "valores": {ym: round(float(v), 2) for ym, v in zip(meses, valores) if v}})
    for g in grupos:
        for i, ym in enumerate(meses):
            soma = round(sum(it["valores"].get(ym, 0.0) for it in g["itens"]), 2)
            total = totais.get(g["grupo"], [None] * len(meses))[i] or 0.0
            if abs(soma - total) > 0.01:
                avisos.append(f"{g['grupo']} em {ym}: os itens somam {soma:,.2f} e o total da planilha diz {total:,.2f}.")
    for i, ym in enumerate(meses):
        soma = round(sum((totais.get(g["grupo"], [0] * len(meses))[i] or 0.0) for g in grupos), 2)
        geral = totais.get("_geral", [0] * len(meses))[i] or 0.0
        if abs(soma - geral) > 0.01:
            avisos.append(f"Total das despesas em {ym}: os grupos somam {soma:,.2f} e a planilha diz {geral:,.2f}.")
    return {"meses": meses, "grupos": grupos, "avisos": avisos}


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    base = ler(Path(sys.argv[1]))
    for a in base["avisos"]:
        print(">>> AVISO:", a)
    atual = json.loads(SAIDA.read_text(encoding="utf-8")) if SAIDA.exists() else {}
    atual.update({"origem": Path(sys.argv[1]).name, "meses": base["meses"], "grupos": base["grupos"]})
    SAIDA.write_text(json.dumps(atual, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    n = sum(len(g["itens"]) for g in base["grupos"])
    print(f">>> {len(base['grupos'])} grupos, {n} itens, {len(base['meses'])} meses gravados em {SAIDA.name}")


if __name__ == "__main__":
    main()
